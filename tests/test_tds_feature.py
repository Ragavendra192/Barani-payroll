import unittest
import openpyxl
from services.excel_service import generate_attendance_template_excel
from services.wages_excel_exporter import generate_wages_excel
from models.employee import get_all_employees

class TestTDSFeature(unittest.TestCase):
    def test_attendance_template_tds_column(self):
        employees = get_all_employees(status='Active')
        bio = generate_attendance_template_excel(2026, 7, employees)
        wb = openpyxl.load_workbook(bio)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn('LIC', headers)
        self.assertIn('TDS', headers)
        lic_idx = headers.index('LIC')
        tds_idx = headers.index('TDS')
        self.assertEqual(tds_idx, lic_idx + 1, "TDS must be immediately after LIC")

    def test_wages_export_tds_only_in_staff_pf_esi(self):
        bio_wages, fname = generate_wages_excel(2026, 7, 'ALL')
        wb_wages = openpyxl.load_workbook(bio_wages)
        self.assertIn('Staff_PF_ESI', wb_wages.sheetnames)

        for sname in wb_wages.sheetnames:
            ws_s = wb_wages[sname]
            s_headers = [ws_s.cell(4, c).value for c in range(1, ws_s.max_column + 1) if ws_s.cell(4, c).value is not None]
            has_tds = any('TDS' in str(h).upper() for h in s_headers)
            if sname == 'Staff_PF_ESI':
                self.assertTrue(has_tds, "Staff_PF_ESI sheet must have TDS Dedn column")
                lic_idx = next(i for i, h in enumerate(s_headers) if 'LIC' in str(h))
                tds_idx = next(i for i, h in enumerate(s_headers) if 'TDS' in str(h))
                self.assertEqual(tds_idx, lic_idx + 1, "TDS Dedn must be right after LIC Dedn in Staff_PF_ESI")
            else:
                self.assertFalse(has_tds, f"Sheet {sname} must NOT have TDS column")

if __name__ == '__main__':
    unittest.main()
