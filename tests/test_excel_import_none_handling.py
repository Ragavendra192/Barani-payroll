import unittest
import io
import pandas as pd
from app import app

class TestExcelImportNoneHandling(unittest.TestCase):
    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_excel_import_with_nan_and_none_values(self):
        """Test that importing an Excel file with missing/empty cells doesn't crash with float(None)."""
        data = {
            'Emp ID': [1001, 1002],
            'Employee Name': ['KRISHNAN.A.P', 'M.PERUMALSAMY'],
            'Type': ['STAFF', 'STAFF'],
            'Present': [26.0, None],  # One valid, one None/NaN
            'N/H': [None, 1.0],
            'EL': [None, None],
            'CL': [0.0, None],
            'SL': [None, 0.0],
            'OT Hours': [None, None],
            'Opening Advance': [None, None],
            'New Advance': [None, None],
            'Advance Deduction': [None, None],
            'Closing Advance': [None, None],
            'Arrears': [None, None],
            'NAPS': [None, None],
            'LIC': [None, None],
            'Accommodation': [None, None],
            'Other': [None, None]
        }
        df = pd.DataFrame(data)
        excel_io = io.BytesIO()
        df.to_excel(excel_io, index=False)
        excel_io.seek(0)

        response = self.app.post(
            '/attendance?year=2026&month=8',
            data={
                'action': 'import_excel',
                'worker_working_days': '25.5',
                'staff_working_days': '27.0',
                'import_file': (excel_io, 'attendance_test.xlsx')
            },
            content_type='multipart/form-data',
            follow_redirects=True
        )
        self.assertEqual(response.status_code, 200)
        # Should not have error flash about float() argument or Series truth value
        self.assertNotIn(b"float() argument must be a string or a real number, not &#39;NoneType&#39;", response.data)
        self.assertNotIn(b"float() argument must be a string or a real number, not 'NoneType'", response.data)
        self.assertNotIn(b"The truth value of a Series is ambiguous", response.data)
        self.assertIn(b"Monthly data imported and saved successfully", response.data)

    def test_august_worker_wages_excel_export(self):
        """Test importing August attendance with Worker Zaheer Hussain and verifying exported Excel."""
        import openpyxl
        from services.wages_excel_exporter import generate_wages_excel

        data = {
            'Emp ID': [10002],
            'Employee Name': ['D.ZAHEER HUSSAIN'],
            'Type': ['WORKER'],
            'Present': [20.5],
            'N/H': [1.0],
            'EL': [0.0],
            'CL': [1.0],
            'SL': [0.0],
            'OT Hours': [4.5],
            'Opening Advance': [908000.0],
            'New Advance': [0.0],
            'Advance Deduction': [5000.0],
            'Closing Advance': [903000.0],
            'Arrears': [0.0],
            'NAPS': [0.0],
            'LIC': [275.0],
            'Accommodation': [0.0],
            'Other': [0.0]
        }
        df = pd.DataFrame(data)
        excel_io = io.BytesIO()
        df.to_excel(excel_io, index=False)
        excel_io.seek(0)

        # Post import
        self.app.post(
            '/attendance?year=2026&month=8',
            data={
                'action': 'import_excel',
                'worker_working_days': '25.5',
                'staff_working_days': '27.0',
                'import_file': (excel_io, 'aug_import.xlsx')
            },
            content_type='multipart/form-data',
            follow_redirects=True
        )

        # Export wages Excel
        out_excel_io, fname = generate_wages_excel(2026, 8, category_filter='WORKER_PF_ESI')
        self.assertIsNotNone(out_excel_io)
        wb = openpyxl.load_workbook(out_excel_io, data_only=True)
        self.assertIn('Worker_PF_ESI', wb.sheetnames)
        ws = wb['Worker_PF_ESI']

        # Find row for 10002
        zaheer_row = None
        for r in range(5, ws.max_row + 1):
            if str(ws.cell(row=r, column=2).value).strip() in ('10002', 'BHIPL-W1', 'W10002'):
                zaheer_row = r
                break
            # Also check name
            if 'ZAHEER' in str(ws.cell(row=r, column=5).value or ''):
                zaheer_row = r
                break

        if zaheer_row:
            # Check Net Salary
            net_val = float(ws.cell(row=zaheer_row, column=51).value or 0.0) # Column AY / Net Salary
            self.assertAlmostEqual(net_val, 18524.38, delta=5.0)

if __name__ == '__main__':
    unittest.main()
