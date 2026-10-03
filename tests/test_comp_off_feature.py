import unittest
import openpyxl
from decimal import Decimal

from utils.payroll_calculation_engine import calculate_attendance, calculate_payroll
from services.excel_service import generate_attendance_template_excel

class TestCompOffFeature(unittest.TestCase):

    def test_staff_comp_off_in_attendance(self):
        att = {
            'present_days': 20.0,
            'nh': 1.0,
            'el': 1.0,
            'c_off': 2.0,
            'sl': 1.0,
            'cl': 1.0
        }
        res = calculate_attendance(att, is_worker=False, standard_days=27.0)
        self.assertEqual(res['present_days'], 20.0)
        self.assertEqual(res['ph'], 1.0)
        self.assertEqual(res['el'], 1.0)
        self.assertEqual(res['c_off'], 2.0)
        self.assertEqual(res['sl'], 1.0)
        self.assertEqual(res['cl'], 1.0)
        # Total Worked Days = 20 + 1 + 1 + 2 + 1 + 1 = 26.0
        self.assertEqual(res['total_worked_days'], 26.0)

    def test_worker_comp_off_ignored(self):
        att = {
            'present_days': 20.0,
            'nh': 1.0,
            'el': 1.0,
            'c_off': 2.0,  # Worker should NOT receive comp-off
            'sl': 1.0,
            'cl': 1.0
        }
        res = calculate_attendance(att, is_worker=True, standard_days=26.0)
        self.assertEqual(res['present_days'], 20.0)
        self.assertEqual(res['c_off'], 0.0)
        # Total Worked Days = 20 + 1 + 1 + 0 + 1 + 1 = 24.0
        self.assertEqual(res['total_worked_days'], 24.0)

    def test_excel_template_columns_and_formula(self):
        sample_employees = [
            {
                'Employee_ID': 1,
                'Emp_No': '1001',
                'Employee_Name': 'Staff User',
                'Employee_Type': 'STAFF',
                'Category': 'STAFF_PF_ESI'
            },
            {
                'Employee_ID': 2,
                'Emp_No': '101',
                'Employee_Name': 'Worker User',
                'Employee_Type': 'WORKER',
                'Category': 'WORKER_PF_ESI'
            }
        ]
        excel_buf = generate_attendance_template_excel(
            year=2026,
            month=7,
            employees=sample_employees,
            worker_working_days=26.0,
            staff_working_days=27.0
        )
        wb = openpyxl.load_workbook(excel_buf, data_only=False)
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        
        # Verify C-Off is directly after EL
        el_idx = headers.index('EL') + 1
        coff_idx = headers.index('C-Off') + 1
        cl_idx = headers.index('CL') + 1
        sl_idx = headers.index('SL') + 1
        tot_idx = headers.index('Total Present Days') + 1

        self.assertEqual(coff_idx, el_idx + 1, "C-Off must be immediately after EL")
        self.assertEqual(cl_idx, coff_idx + 1, "CL must be immediately after C-Off")
        self.assertEqual(tot_idx, sl_idx + 1, "Total Present Days must be immediately after SL")

        # Verify formula in row 2 (Staff row)
        tot_cell_r2 = ws.cell(row=2, column=tot_idx).value
        # Formula: =Present + NH + EL + C-Off + SL + CL
        self.assertIn('+', tot_cell_r2)
        self.assertTrue(tot_cell_r2.startswith('='))
        
        # Verify Closing Advance formula exists and is valid
        close_adv_idx = headers.index('Closing Advance') + 1
        close_cell_r2 = ws.cell(row=2, column=close_adv_idx).value
        self.assertTrue(close_cell_r2.startswith('=MAX(0,'))

if __name__ == '__main__':
    unittest.main()
