import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import openpyxl
from decimal import Decimal

# Import from our central payroll formula engine
from payroll_formulas import (
    calculate_payroll,
    CAT_STAFF_PF_ESI,
    CAT_WORKER_PF_ESI,
    CAT_STAFF_NAPS,
    CAT_WORKER_NAPS,
    CAT_STAFF_NON_PF_ESI,
    CAT_WORKER_NON_PF_ESI
)

WORKBOOK_PATH = 'e:/payroll_app/samples/BHIPL UNIT - 1 SALARY STATEMENT FOR THE MONTH OF Sep-2026 Final.xlsx'

def num(val, default=0.0):
    if val is None:
        return default
    try:
        s = str(val).strip().replace(',', '')
        if s == '' or s.lower() == 'none' or s.startswith('#'):
            return default
        return float(s)
    except:
        return default

def run_reconciliation():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    print("=" * 80)
    print("BHIPL PAYROLL ENGINE — COMPLETE SEPTEMBER 2026 RECONCILIATION AUDIT")
    print("=" * 80)

    total_employees_audited = 0
    total_perfect_matches = 0
    sheet_stats = {}
    differences_log = []

    # 1. STAFFS (STAFF_PF_ESI)
    ws = wb['STAFFS']
    staff_results = []
    for r in range(6, ws.max_row + 1):
        s_no = ws.cell(r, 2).value
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 7).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        # Excel values
        xl_fixed_gross = num(ws.cell(r, 26).value) # Col Z
        xl_present = num(ws.cell(r, 14).value)      # Col N
        xl_nh = num(ws.cell(r, 15).value)           # Col O
        xl_coff = num(ws.cell(r, 16).value)         # Col P
        xl_el = num(ws.cell(r, 17).value)           # Col Q
        xl_cl = num(ws.cell(r, 18).value)           # Col R
        xl_sl = num(ws.cell(r, 19).value)           # Col S
        xl_total_days = num(ws.cell(r, 20).value)   # Col T

        xl_earned_gross = num(ws.cell(r, 32).value) # Col AF
        xl_pf_gross = num(ws.cell(r, 33).value)     # Col AG
        xl_esi_gross = num(ws.cell(r, 34).value)    # Col AH
        xl_pf_ded = num(ws.cell(r, 35).value)       # Col AI
        xl_esi_ded = num(ws.cell(r, 36).value)      # Col AJ
        xl_nats = num(ws.cell(r, 37).value)         # Col AK
        xl_lic = num(ws.cell(r, 38).value)          # Col AL
        xl_tds = num(ws.cell(r, 39).value)          # Col AM
        xl_adv = num(ws.cell(r, 40).value)          # Col AN
        xl_tot_ded = num(ws.cell(r, 41).value)      # Col AO
        xl_net = num(ws.cell(r, 42).value)          # Col AP

        xl_new_adv = num(ws.cell(r, 44).value)      # Col AR
        xl_inst = num(ws.cell(r, 45).value)         # Col AS
        xl_op_adv = num(ws.cell(r, 46).value)       # Col AT
        xl_cl_adv = num(ws.cell(r, 47).value)       # Col AU

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_STAFF_PF_ESI, 'Employee_Type': 'STAFF'}
        sal = {'Fixed_Gross': xl_fixed_gross}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'c_off': xl_coff,
            'el': xl_el, 'cl': xl_cl, 'sl': xl_sl, 'total_days': xl_total_days
        }
        ded = {
            'lic': xl_lic, 'tds': xl_tds, 'advance': xl_adv, 'naps': xl_nats,
            'new_adv': xl_new_adv, 'installment': xl_inst, 'opening_adv': xl_op_adv
        }

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        # Comparisons
        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['PF_Gross'] - xl_pf_gross) > 0.05:
            diffs.append(f"PF Gross (calc {calc['PF_Gross']:.2f} vs xl {xl_pf_gross:.2f})")
        if abs(calc['ESI_Gross'] - xl_esi_gross) > 0.05:
            diffs.append(f"ESI Gross (calc {calc['ESI_Gross']:.2f} vs xl {xl_esi_gross:.2f})")
        if abs(calc['PF_Deduction'] - xl_pf_ded) > 0.05:
            diffs.append(f"PF Ded (calc {calc['PF_Deduction']:.2f} vs xl {xl_pf_ded:.2f})")
        if abs(calc['ESI_Deduction'] - xl_esi_ded) > 0.05:
            diffs.append(f"ESI Ded (calc {calc['ESI_Deduction']:.2f} vs xl {xl_esi_ded:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        staff_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_STAFF_PF_ESI, emp_no_str, name, diffs))

    sheet_stats['STAFF_PF_ESI'] = {
        'total': len(staff_results),
        'matched': sum(1 for _, _, d in staff_results if not d),
        'rate': (sum(1 for _, _, d in staff_results if not d) / len(staff_results) * 100) if staff_results else 100
    }

    # 2. WORKER'S - ESI PF (WORKER_PF_ESI)
    ws = wb["WORKER'S - ESI PF"]
    worker_pf_results = []
    for r in range(5, ws.max_row + 1):
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 8).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        xl_pdw = num(ws.cell(r, 27).value)        # Col AA
        xl_fixed_gross = num(ws.cell(r, 34).value) # Col AH
        xl_present = num(ws.cell(r, 16).value)    # Col P
        xl_nh = num(ws.cell(r, 17).value)         # Col Q
        xl_el = num(ws.cell(r, 18).value)         # Col R
        xl_cl = num(ws.cell(r, 19).value)         # Col S
        xl_sl = num(ws.cell(r, 20).value)         # Col T
        xl_total_days = num(ws.cell(r, 21).value) # Col U
        xl_act_ot = num(ws.cell(r, 22).value)     # Col V

        xl_ot_wages = num(ws.cell(r, 41).value)   # Col AO
        xl_earned_gross = num(ws.cell(r, 42).value) # Col AP
        xl_pf_gross = num(ws.cell(r, 44).value)   # Col AR
        xl_esi_gross = num(ws.cell(r, 45).value)  # Col AS
        xl_arrears = num(ws.cell(r, 46).value)    # Col AT
        xl_pf_ded_cell = ws.cell(r, 47).value     # Col AU
        xl_pf_ded = num(xl_pf_ded_cell)
        xl_esi_ded = num(ws.cell(r, 48).value)    # Col AV
        xl_lic = num(ws.cell(r, 49).value)        # Col AW
        xl_adv = num(ws.cell(r, 50).value)        # Col AX
        xl_naps = num(ws.cell(r, 51).value)       # Col AY
        xl_accom = num(ws.cell(r, 52).value)      # Col AZ
        xl_tot_ded = num(ws.cell(r, 53).value)    # Col BA
        xl_net = num(ws.cell(r, 54).value)        # Col BB

        pf_eligible = True
        if xl_pf_ded_cell is None or str(xl_pf_ded_cell).strip() == '':
            pf_eligible = False

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_WORKER_PF_ESI, 'Employee_Type': 'WORKER'}
        sal = {'Per_Day_Wage': xl_pdw, 'Fixed_Gross': xl_fixed_gross, 'PF_Eligible': pf_eligible}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'el': xl_el, 'cl': xl_cl, 'sl': xl_sl,
            'total_days': xl_total_days, 'actual_ot_hours': xl_act_ot
        }
        ded = {'lic': xl_lic, 'advance': xl_adv, 'naps': xl_naps, 'accommodation': xl_accom, 'arrears': xl_arrears}

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['OT_Wages'] - xl_ot_wages) > 0.05:
            diffs.append(f"OT Wages (calc {calc['OT_Wages']:.2f} vs xl {xl_ot_wages:.2f})")
        if abs(calc['PF_Gross'] - xl_pf_gross) > 0.05:
            diffs.append(f"PF Gross (calc {calc['PF_Gross']:.2f} vs xl {xl_pf_gross:.2f})")
        if abs(calc['ESI_Gross'] - xl_esi_gross) > 0.05:
            diffs.append(f"ESI Gross (calc {calc['ESI_Gross']:.2f} vs xl {xl_esi_gross:.2f})")
        if abs(calc['PF_Deduction'] - xl_pf_ded) > 0.05:
            diffs.append(f"PF Ded (calc {calc['PF_Deduction']:.2f} vs xl {xl_pf_ded:.2f})")
        if abs(calc['ESI_Deduction'] - xl_esi_ded) > 0.05:
            diffs.append(f"ESI Ded (calc {calc['ESI_Deduction']:.2f} vs xl {xl_esi_ded:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        worker_pf_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_WORKER_PF_ESI, emp_no_str, name, diffs))

    sheet_stats['WORKER_PF_ESI'] = {
        'total': len(worker_pf_results),
        'matched': sum(1 for _, _, d in worker_pf_results if not d),
        'rate': (sum(1 for _, _, d in worker_pf_results if not d) / len(worker_pf_results) * 100) if worker_pf_results else 100
    }

    # 3. STAFF'S (NAPS) (STAFF_NAPS)
    ws = wb["STAFF'S (NAPS)"]
    staff_naps_results = []
    for r in range(5, ws.max_row + 1):
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 6).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        xl_fixed_gross = num(ws.cell(r, 31).value) # Col AE
        xl_present = num(ws.cell(r, 13).value)     # Col M
        xl_nh = num(ws.cell(r, 14).value)          # Col N
        xl_coff = num(ws.cell(r, 15).value)        # Col O
        xl_el = num(ws.cell(r, 16).value)          # Col P
        xl_cl = num(ws.cell(r, 17).value)          # Col Q
        xl_sl = num(ws.cell(r, 18).value)          # Col R
        xl_total_days = num(ws.cell(r, 19).value)  # Col S

        xl_earned_gross = num(ws.cell(r, 39).value) # Col AM
        xl_arrears = num(ws.cell(r, 40).value)      # Col AN
        xl_naps_ded = num(ws.cell(r, 41).value)     # Col AO
        xl_lic = num(ws.cell(r, 42).value)          # Col AP
        xl_adv = num(ws.cell(r, 43).value)          # Col AQ
        xl_accom = num(ws.cell(r, 44).value)        # Col AR
        xl_tot_ded = num(ws.cell(r, 45).value)      # Col AS
        xl_net = num(ws.cell(r, 46).value)          # Col AT

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_STAFF_NAPS, 'Employee_Type': 'STAFF'}
        sal = {'Fixed_Gross': xl_fixed_gross}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'c_off': xl_coff,
            'el': xl_el, 'cl': xl_cl, 'sl': xl_sl, 'total_days': xl_total_days
        }
        ded = {'lic': xl_lic, 'advance': xl_adv, 'accommodation': xl_accom, 'naps': xl_naps_ded, 'arrears': xl_arrears}

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['NAPS_Deduction'] - xl_naps_ded) > 0.05:
            diffs.append(f"NAPS Ded (calc {calc['NAPS_Deduction']:.2f} vs xl {xl_naps_ded:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        staff_naps_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_STAFF_NAPS, emp_no_str, name, diffs))

    sheet_stats['STAFF_NAPS'] = {
        'total': len(staff_naps_results),
        'matched': sum(1 for _, _, d in staff_naps_results if not d),
        'rate': (sum(1 for _, _, d in staff_naps_results if not d) / len(staff_naps_results) * 100) if staff_naps_results else 100
    }

    # 4. WORKER'S (NAPS) (2) (WORKER_NAPS)
    ws = wb["WORKER'S (NAPS) (2)"]
    worker_naps_results = []
    for r in range(5, ws.max_row + 1):
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 6).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        xl_pdw = num(ws.cell(r, 24).value)         # Col X
        xl_fixed_gross = num(ws.cell(r, 31).value) # Col AE
        xl_present = num(ws.cell(r, 13).value)     # Col M
        xl_nh = num(ws.cell(r, 14).value)          # Col N
        xl_el = num(ws.cell(r, 15).value)          # Col O
        xl_cl = num(ws.cell(r, 16).value)          # Col P
        xl_sl = num(ws.cell(r, 17).value)          # Col Q
        xl_total_days = num(ws.cell(r, 18).value)  # Col R
        xl_act_ot = num(ws.cell(r, 19).value)      # Col S

        xl_ot_wages = num(ws.cell(r, 38).value)    # Col AL
        xl_earned_gross = num(ws.cell(r, 39).value) # Col AM
        xl_arrears = num(ws.cell(r, 40).value)      # Col AN
        xl_naps_ded = num(ws.cell(r, 41).value)     # Col AO
        xl_lic = num(ws.cell(r, 42).value)          # Col AP
        xl_adv = num(ws.cell(r, 43).value)          # Col AQ
        xl_accom = num(ws.cell(r, 44).value)        # Col AR
        xl_tot_ded = num(ws.cell(r, 45).value)      # Col AS
        xl_net = num(ws.cell(r, 46).value)          # Col AT

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_WORKER_NAPS, 'Employee_Type': 'WORKER'}
        sal = {'Per_Day_Wage': xl_pdw, 'Fixed_Gross': xl_fixed_gross}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'el': xl_el, 'cl': xl_cl, 'sl': xl_sl,
            'total_days': xl_total_days, 'actual_ot_hours': xl_act_ot
        }
        ded = {'lic': xl_lic, 'advance': xl_adv, 'accommodation': xl_accom, 'naps': xl_naps_ded, 'arrears': xl_arrears}

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['OT_Wages'] - xl_ot_wages) > 0.05:
            diffs.append(f"OT Wages (calc {calc['OT_Wages']:.2f} vs xl {xl_ot_wages:.2f})")
        if abs(calc['NAPS_Deduction'] - xl_naps_ded) > 0.05:
            diffs.append(f"NAPS Ded (calc {calc['NAPS_Deduction']:.2f} vs xl {xl_naps_ded:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        worker_naps_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_WORKER_NAPS, emp_no_str, name, diffs))

    sheet_stats['WORKER_NAPS'] = {
        'total': len(worker_naps_results),
        'matched': sum(1 for _, _, d in worker_naps_results if not d),
        'rate': (sum(1 for _, _, d in worker_naps_results if not d) / len(worker_naps_results) * 100) if worker_naps_results else 100
    }

    # 5. Non-pf ESi staff (STAFF_NON_PF_ESI)
    ws = wb['Non-pf ESi staff']
    staff_non_pf_results = []
    for r in range(5, ws.max_row + 1):
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 6).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        xl_fixed_gross = num(ws.cell(r, 31).value) # Col AE
        xl_present = num(ws.cell(r, 13).value)     # Col M
        xl_nh = num(ws.cell(r, 14).value)          # Col N
        xl_coff = num(ws.cell(r, 15).value)        # Col O
        xl_el = num(ws.cell(r, 16).value)          # Col P
        xl_cl = num(ws.cell(r, 17).value)          # Col Q
        xl_sl = num(ws.cell(r, 18).value)          # Col R
        xl_total_days = num(ws.cell(r, 19).value)  # Col S

        xl_earned_gross = num(ws.cell(r, 39).value) # Col AM
        xl_arrears = num(ws.cell(r, 40).value)      # Col AN
        xl_lic = num(ws.cell(r, 42).value)          # Col AP
        xl_adv = num(ws.cell(r, 43).value)          # Col AQ
        xl_accom = num(ws.cell(r, 44).value)        # Col AR
        xl_tot_ded = num(ws.cell(r, 45).value)      # Col AS
        xl_net = num(ws.cell(r, 46).value)          # Col AT

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_STAFF_NON_PF_ESI, 'Employee_Type': 'STAFF'}
        sal = {'Fixed_Gross': xl_fixed_gross}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'c_off': xl_coff,
            'el': xl_el, 'cl': xl_cl, 'sl': xl_sl, 'total_days': xl_total_days
        }
        ded = {'lic': xl_lic, 'advance': xl_adv, 'accommodation': xl_accom, 'arrears': xl_arrears}

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        staff_non_pf_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_STAFF_NON_PF_ESI, emp_no_str, name, diffs))

    sheet_stats['STAFF_NON_PF_ESI'] = {
        'total': len(staff_non_pf_results),
        'matched': sum(1 for _, _, d in staff_non_pf_results if not d),
        'rate': (sum(1 for _, _, d in staff_non_pf_results if not d) / len(staff_non_pf_results) * 100) if staff_non_pf_results else 100
    }

    # 6. Non-pf ESi worker (WORKER_NON_PF_ESI)
    ws = wb['Non-pf ESi worker']
    worker_non_pf_results = []
    for r in range(5, ws.max_row + 1):
        emp_no = ws.cell(r, 4).value
        name = ws.cell(r, 6).value
        if not emp_no or str(emp_no).strip().lower() in ('none', '', 'total', 'grand total'):
            continue
        emp_no_str = str(emp_no).strip()

        xl_pdw = num(ws.cell(r, 24).value)         # Col X
        xl_fixed_gross = num(ws.cell(r, 31).value) # Col AE
        xl_present = num(ws.cell(r, 13).value)     # Col M
        xl_nh = num(ws.cell(r, 14).value)          # Col N
        xl_el = num(ws.cell(r, 15).value)          # Col O
        xl_cl = num(ws.cell(r, 16).value)          # Col P
        xl_sl = num(ws.cell(r, 17).value)          # Col Q
        xl_total_days = num(ws.cell(r, 18).value)  # Col R
        xl_act_ot = num(ws.cell(r, 19).value)      # Col S

        xl_ot_wages = num(ws.cell(r, 38).value)    # Col AL
        xl_earned_gross = num(ws.cell(r, 39).value) # Col AM
        xl_arrears = num(ws.cell(r, 40).value)      # Col AN
        xl_naps = num(ws.cell(r, 41).value)         # Col AO
        xl_lic = num(ws.cell(r, 42).value)          # Col AP
        xl_adv = num(ws.cell(r, 43).value)          # Col AQ
        xl_accom = num(ws.cell(r, 44).value)        # Col AR
        xl_tot_ded = num(ws.cell(r, 45).value)      # Col AS
        xl_net = num(ws.cell(r, 46).value)          # Col AT

        emp = {'Emp_No': emp_no_str, 'Employee_Name': str(name), 'Category': CAT_WORKER_NON_PF_ESI, 'Employee_Type': 'WORKER'}
        sal = {'Per_Day_Wage': xl_pdw, 'Fixed_Gross': xl_fixed_gross}
        att = {
            'present_days': xl_present, 'nh': xl_nh, 'el': xl_el, 'cl': xl_cl, 'sl': xl_sl,
            'total_days': xl_total_days, 'actual_ot_hours': xl_act_ot
        }
        ded = {'lic': xl_lic, 'advance': xl_adv, 'naps': xl_naps, 'accommodation': xl_accom, 'arrears': xl_arrears}

        calc = calculate_payroll(emp, sal, att, ded, standard_days=26.0)

        diffs = []
        if abs(calc['Gross_Wages'] - xl_earned_gross) > 0.05:
            diffs.append(f"Gross (calc {calc['Gross_Wages']:.2f} vs xl {xl_earned_gross:.2f})")
        if abs(calc['OT_Wages'] - xl_ot_wages) > 0.05:
            diffs.append(f"OT Wages (calc {calc['OT_Wages']:.2f} vs xl {xl_ot_wages:.2f})")
        if abs(calc['Total_Deduction'] - xl_tot_ded) > 0.05:
            diffs.append(f"Tot Ded (calc {calc['Total_Deduction']:.2f} vs xl {xl_tot_ded:.2f})")
        if abs(calc['Net_Salary'] - xl_net) > 0.05:
            diffs.append(f"Net Sal (calc {calc['Net_Salary']:.2f} vs xl {xl_net:.2f})")

        worker_non_pf_results.append((emp_no_str, name, diffs))
        total_employees_audited += 1
        if not diffs:
            total_perfect_matches += 1
        else:
            differences_log.append((CAT_WORKER_NON_PF_ESI, emp_no_str, name, diffs))

    sheet_stats['WORKER_NON_PF_ESI'] = {
        'total': len(worker_non_pf_results),
        'matched': sum(1 for _, _, d in worker_non_pf_results if not d),
        'rate': (sum(1 for _, _, d in worker_non_pf_results if not d) / len(worker_non_pf_results) * 100) if worker_non_pf_results else 100
    }

    # Summary Output
    print(f"\nAUDIT RESULTS BY CATEGORY:")
    print("-" * 80)
    for cat, stat in sheet_stats.items():
        print(f"Category: {cat:<22} | Total: {stat['total']:>3} | Matched: {stat['matched']:>3} | Accuracy: {stat['rate']:>6.2f}%")
    print("-" * 80)
    overall_rate = (total_perfect_matches / total_employees_audited * 100) if total_employees_audited else 100
    print(f"OVERALL RECONCILIATION ACCURACY: {total_perfect_matches}/{total_employees_audited} ({overall_rate:.2f}%)\n")

    if differences_log:
        print(f"IDENTIFIED DISCREPANCIES ({len(differences_log)} items):")
        print("-" * 80)
        for cat, emp_no, name, diffs in differences_log:
            print(f"[{cat}] Emp #{emp_no} ({name}): {'; '.join(diffs)}")
    else:
        print("ALL EMPLOYEE PAYROLL CALCULATIONS MATCH THE MANUAL WORKBOOK EXACTLY (100% RECONCILED)!")

if __name__ == '__main__':
    run_reconciliation()
