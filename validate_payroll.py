import os
import openpyxl
from decimal import Decimal, ROUND_HALF_UP
from utils.payroll_calculation_engine import calculate_payroll, money

EXCEL_PATH = r"C:\Users\Ragzz\Downloads\BHIPL UNIT - 1 SALARY STATEMENT FOR THE MONTH OF July-2026.xlsx"

def d(val):
    if val is None or str(val).strip() == '':
        return Decimal('0.00')
    try:
        return Decimal(str(val)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except:
        return Decimal('0.00')

def run_july_2026_validation():
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: Excel workbook not found at {EXCEL_PATH}")
        return

    wb_values = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    total_tested = 0
    passed_count = 0
    failed_count = 0
    failures = []

    sheet_category_map = [
        ("STAFFS", "STAFF_PF_ESI", 27.0, {
            'emp_code_col': 3, 'name_col': 7, 'total_days_col': 20, 'gross_fixed_col': 26,
            'earned_gross_col': 32, 'pf_ded_col': 35, 'esi_ded_col': 37, 'lic_col': 40,
            'advance_col': 42, 'total_ded_col': 43, 'net_sal_col': 44
        }),
        ("WORKER'S - ESI PF", "WORKER_PF_ESI", 26.0, {
            'emp_code_col': 4, 'name_col': 8, 'total_days_col': 21, 'act_ot_col': 22,
            'per_day_wage_col': 26, 'earned_gross_col': 41, 'pf_ded_col': 46,
            'esi_ded_col': 48, 'lic_col': 50, 'advance_col': 51, 'total_ded_col': 54, 'net_sal_col': 55
        }),
        ("STAFF'S (NAPS)", "STAFF_NAPS", 26.0, {
            'emp_code_col': 4, 'name_col': 6, 'total_days_col': 18, 'gross_fixed_col': 30,
            'earned_gross_col': 38, 'naps_col': 40, 'advance_col': 42, 'total_ded_col': 44, 'net_sal_col': 45
        }),
        ("WORKER'S (NAPS) (2)", "WORKER_NAPS", 27.0, {
            'emp_code_col': 4, 'name_col': 6, 'total_days_col': 18, 'act_ot_col': 19,
            'per_day_wage_col': 23, 'earned_gross_col': 38, 'naps_col': 40,
            'advance_col': 42, 'total_ded_col': 44, 'net_sal_col': 45
        }),
        ("Non-pf ESi staff", "STAFF_NON_PF_ESI", 27.0, {
            'emp_code_col': 4, 'name_col': 6, 'total_days_col': 18, 'gross_fixed_col': 30,
            'earned_gross_col': 38, 'advance_col': 42, 'total_ded_col': 44, 'net_sal_col': 45
        }),
        ("Non-pf ESi worker", "WORKER_NON_PF_ESI", 27.0, {
            'emp_code_col': 4, 'name_col': 6, 'total_days_col': 18, 'act_ot_col': 19,
            'per_day_wage_col': 23, 'earned_gross_col': 38, 'advance_col': 42,
            'total_ded_col': 44, 'net_sal_col': 45
        })
    ]

    print("\n" + "="*90)
    print(" BHIPL JULY-2026 SALARY STATEMENT — EMPIRICAL PAYROLL VALIDATION")
    print("="*90)
    print(f"{'Emp Code':<10} | {'Category':<18} | {'Field':<18} | {'Excel Val':<12} | {'Engine Val':<12} | {'Diff':<8} | {'Status'}")
    print("-" * 90)

    for sheet_name, category, std_days, col_map in sheet_category_map:
        ws = wb_values[sheet_name]
        
        for r in range(5, ws.max_row + 1):
            emp_code = ws.cell(r, col_map['emp_code_col']).value
            emp_name = ws.cell(r, col_map['name_col']).value

            if emp_code is None or emp_name is None:
                continue
                
            name_str = str(emp_name).strip()
            emp_code_str = str(emp_code).strip()

            if name_str == '' or emp_code_str == '' or name_str.upper().startswith("TOTAL") or name_str.upper().startswith("GRAND"):
                continue

            total_days = d(ws.cell(r, col_map['total_days_col']).value)
            act_ot = d(ws.cell(r, col_map['act_ot_col']).value) if 'act_ot_col' in col_map else Decimal('0.00')

            gross_fixed = d(ws.cell(r, col_map['gross_fixed_col']).value) if 'gross_fixed_col' in col_map else Decimal('0.00')
            per_day_wage = d(ws.cell(r, col_map['per_day_wage_col']).value) if 'per_day_wage_col' in col_map else Decimal('0.00')

            expected_gross = d(ws.cell(r, col_map['earned_gross_col']).value)
            expected_pf = d(ws.cell(r, col_map['pf_ded_col']).value) if 'pf_ded_col' in col_map else Decimal('0.00')
            expected_esi = d(ws.cell(r, col_map['esi_ded_col']).value) if 'esi_ded_col' in col_map else Decimal('0.00')
            expected_lic = d(ws.cell(r, col_map['lic_col']).value) if 'lic_col' in col_map else Decimal('0.00')
            expected_naps = d(ws.cell(r, col_map['naps_col']).value) if 'naps_col' in col_map else Decimal('0.00')
            expected_advance = d(ws.cell(r, col_map['advance_col']).value) if 'advance_col' in col_map else Decimal('0.00')
            expected_tot_ded = d(ws.cell(r, col_map['total_ded_col']).value)
            expected_net = d(ws.cell(r, col_map['net_sal_col']).value)

            emp_dict = {
                'Employee_ID': emp_code_str,
                'Emp_No': emp_code_str,
                'Name': name_str,
                'Category': category,
                'Employee_Type': 'STAFF' if 'STAFF' in category else 'WORKER',
                'Payroll_Category': 'PF_ESI' if 'PF_ESI' in category else ('NAPS' if 'NAPS' in category else 'NON_PF_ESI')
            }

            sal_dict = {
                'Gross_Wages': float(gross_fixed),
                'Base_Gross': float(gross_fixed),
                'Per_Day_Wage': float(per_day_wage),
                'PF_Eligible': 'PF_ESI' in category,
                'ESI_Eligible': 'PF_ESI' in category
            }

            att_dict = {
                'present_days': float(total_days),
                'total_days': float(total_days),
                'actual_ot_hours': float(act_ot)
            }

            ded_dict = {
                'lic': float(expected_lic),
                'naps': float(expected_naps),
                'advance': float(expected_advance)
            }

            calc_res = calculate_payroll(emp_dict, sal_dict, att_dict, ded_dict, standard_days=std_days)

            engine_gross = d(calc_res['Gross_Wages'])
            engine_pf = d(calc_res['PF_Deduction'])
            engine_esi = d(calc_res['ESI_Deduction'])
            engine_tot_ded = d(calc_res['Total_Deduction'])
            engine_net = d(calc_res['Net_Salary'])

            fields_to_check = [
                ('Earned Gross', expected_gross, engine_gross),
                ('PF Deduction', expected_pf, engine_pf),
                ('ESI Deduction', expected_esi, engine_esi),
                ('Total Deduction', expected_tot_ded, engine_tot_ded),
                ('Net Pay', expected_net, engine_net)
            ]

            emp_has_failure = False
            total_tested += 1

            for f_name, exp_val, eng_val in fields_to_check:
                diff = eng_val - exp_val
                status = "PASS" if abs(diff) <= Decimal('0.05') else "FAIL"
                if status == "FAIL":
                    emp_has_failure = True
                    failures.append((emp_code_str, name_str, category, f_name, exp_val, eng_val, diff))
                    print(f"{emp_code_str:<10} | {category:<18} | {f_name:<18} | {exp_val:<12} | {eng_val:<12} | {diff:<8} | FAIL")

            if not emp_has_failure:
                passed_count += 1
                print(f"{emp_code_str:<10} | {category:<18} | {'Net Pay':<18} | {expected_net:<12} | {engine_net:<12} | {0.00:<8} | PASS")

    failed_count = total_tested - passed_count

    print("\n" + "="*90)
    print(" VALIDATION SUMMARY RESULTS")
    print("="*90)
    print(f"Total employees tested: {total_tested}")
    print(f"PASS:                   {passed_count}")
    print(f"FAIL:                   {failed_count}")
    print("="*90)

    if failures:
        print("\nDETAILED FAILURES LIST:")
        for f in failures:
            print(f"  Emp: {f[0]} ({f[1]}) | Category: {f[2]} | Field: {f[3]} | Expected: {f[4]} | Actual: {f[5]} | Diff: {f[6]}")

if __name__ == '__main__':
    run_july_2026_validation()
