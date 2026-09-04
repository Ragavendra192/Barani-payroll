import io
import pandas as pd
from models.payroll_transaction import get_payroll_transactions

def generate_monthly_salary_statement_excel(year, month):
    """
    Generates a 6-sheet Excel Workbook reproducing the exact BHIPL Unit-I Salary Statement format:
      1. STAFFS (STAFF_PF_ESI)
      2. WORKER'S - ESI PF (WORKER_PF_ESI)
      3. STAFF'S (NAPS) (STAFF_NAPS)
      4. WORKER'S (NAPS) (WORKER_NAPS)
      5. Non-pf ESi staff (STAFF_NON_PF_ESI)
      6. Non-pf ESi worker (WORKER_NON_PF_ESI)
    """
    output = io.BytesIO()
    
    category_sheets = [
        ("STAFFS", "STAFF_PF_ESI"),
        ("WORKER'S - ESI PF", "WORKER_PF_ESI"),
        ("STAFF'S (NAPS)", "STAFF_NAPS"),
        ("WORKER'S (NAPS)", "WORKER_NAPS"),
        ("Non-pf ESi staff", "STAFF_NON_PF_ESI"),
        ("Non-pf ESi worker", "WORKER_NON_PF_ESI")
    ]

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet_title, cat in category_sheets:
            records = get_payroll_transactions(year, month, category=cat)
            
            if not records:
                # Add empty dataframe with standard columns
                df = pd.DataFrame(columns=[
                    "Emp No", "ERP Emp No", "Name", "Department", "Designation",
                    "Working Days", "OT Hours", "Per Day Wage",
                    "Basic + DA", "HRA", "Conveyance", "Washing", "Other",
                    "Special Allowance", "OT Wages", "Gross Wages",
                    "PF", "ESI", "NAPS", "LIC", "Opening Advance", "New Advance", "Advance", "Closing Advance", "Accommodation", "Other Dedn", "Total Dedn", "Net Salary"
                ])
            else:
                formatted_rows = []
                for r in records:
                    formatted_rows.append({
                        "Emp No": r.get('Emp_No'),
                        "ERP Emp No": r.get('ERP_Emp_No') or r.get('Emp_No'),
                        "Emp Code": r.get('Emp_Code'),
                        "Name": r.get('Employee_Name'),
                        "Department": r.get('Department'),
                        "Designation": r.get('Designation'),
                        "Grade": r.get('Grade'),
                        "Present Days": r.get('Present_Days', 0.0),
                        "Total Days": r.get('Total_Days', 0.0),
                        "Working Days": r.get('Working_Days', 0.0),
                        "OT Hours": r.get('OT_Hours', 0.0),
                        "Per Day Wage": r.get('Per_Day_Wage', 0.0),
                        "Basic + DA": r.get('Basic_DA_Earned', 0.0),
                        "HRA": r.get('HRA_Earned', 0.0),
                        "Conveyance": r.get('Conveyance_Earned', 0.0),
                        "Washing": r.get('Washing_Allowance_Earned', 0.0),
                        "Other": r.get('Other_Allowance_Earned', 0.0),
                        "Special Allowance": r.get('Special_Allowance_Earned', 0.0),
                        "OT Wages": r.get('OT_Wages', 0.0),
                        "Gross Wages": r.get('Gross_Wages', 0.0),
                        "PF": r.get('PF_Deduction', 0.0),
                        "ESI": r.get('ESI_Deduction', 0.0),
                        "NAPS": r.get('NAPS_Deduction', 0.0),
                        "LIC": r.get('LIC_Deduction', 0.0),
                        "Opening Advance": r.get('Opening_Advance', 0.0),
                        "New Advance": r.get('New_Advance', 0.0),
                        "Advance": r.get('Advance_Deduction', 0.0),
                        "Closing Advance": r.get('Closing_Advance', 0.0),
                        "Accommodation": r.get('Accommodation_Deduction', 0.0),
                        "Arrears": r.get('Arrears', 0.0),
                        "Total Dedn": r.get('Total_Deduction', 0.0),
                        "Net Salary": r.get('Net_Salary', 0.0)
                    })
                df = pd.DataFrame(formatted_rows)
            
            df.to_excel(writer, sheet_name=sheet_title, index=False)

    output.seek(0)
    return output


def generate_employee_master_excel(employees=None, include_sample=False):
    """
    Generate professional Employee Master Excel workbook pre-filled with live data
    ready for HR to edit and re-upload.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if employees is None:
        from models.employee import get_all_employees
        employees = get_all_employees(status=None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Employee_Master'
    ws.views.sheetView[0].showGridLines = True

    # Columns configuration: (Header name, Dict Key, Type, Default width)
    columns_config = [
        ('Emp_No', 'Emp_No', 'text', 12),
        ('Employee_Name', 'Employee_Name', 'text', 24),
        ('Employee_Type', 'Employee_Type', 'center', 16),
        ('Payroll_Category', 'Payroll_Category', 'center', 18),
        ('Department', 'Department', 'text', 18),
        ('Designation', 'Designation', 'text', 20),
        ('Grade', 'Grade', 'center', 12),
        ('Status', 'Status', 'center', 12),
        ('DOJ', 'DOJ', 'center', 14),
        ('DOB', 'DOB', 'center', 14),
        ('Father_Name', 'Father_Name', 'text', 22),
        ('Bank_Acc_No', 'Bank_Acc_No', 'text', 20),
        ('Bank_IFSC', 'Bank_IFSC', 'text', 15),
        ('UAN_No', 'UAN_No', 'text', 18),
        ('ESI_No', 'ESI_No', 'text', 18),
        ('Phone_Number', 'Phone_Number', 'text', 16),
        ('Email_ID', 'Email_ID', 'text', 26),
        ('Fixed_Gross', 'Fixed_Gross', 'currency', 14),
        ('Basic_DA', 'Basic_DA', 'currency', 14),
        ('HRA', 'HRA', 'currency', 12),
        ('Conveyance_Allowance', 'Conveyance_Allowance', 'currency', 20),
        ('Washing_Allowance', 'Washing_Allowance', 'currency', 18),
        ('Other_Allowance', 'Other_Allowance', 'currency', 18),
        ('Per_Day_Wage', 'Per_Day_Wage', 'currency', 15),
        ('OT_Rate', 'OT_Rate', 'currency', 12),
        ('PF_Eligible', 'PF_Eligible', 'center', 14),
        ('ESI_Eligible', 'ESI_Eligible', 'center', 14),
        ('LIC', 'LIC', 'currency', 12),
    ]

    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')  # Deep Navy Blue
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border_side = Side(style='thin', color='D1D5DB')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    regular_font = Font(name='Segoe UI', size=10)

    # Write headers
    headers = [c[0] for c in columns_config]
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(columns_config) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = cell_border

    def clean_str_val(v):
        if v is None or pd.isna(v):
            return ''
        s = str(v).strip()
        if s.endswith('.0') and s[:-2].isdigit():
            return s[:-2]
        if s.lower() in ['nan', 'none', 'null']:
            return ''
        return s

    rows_to_write = []
    if employees:
        for emp in employees:
            pf_val = 'YES' if emp.get('PF_Eligible') in [True, 1, '1', 'YES', 'True'] else 'NO'
            esi_val = 'YES' if emp.get('ESI_Eligible') in [True, 1, '1', 'YES', 'True'] else 'NO'

            row_data = [
                clean_str_val(emp.get('Emp_No')),
                clean_str_val(emp.get('Employee_Name') or emp.get('Emp_Name')),
                clean_str_val(emp.get('Employee_Type') or 'STAFF'),
                clean_str_val(emp.get('Payroll_Category') or 'PF_ESI'),
                clean_str_val(emp.get('Department')),
                clean_str_val(emp.get('Designation')),
                clean_str_val(emp.get('Grade')),
                clean_str_val(emp.get('Status') or 'Active'),
                clean_str_val(emp.get('DOJ')),
                clean_str_val(emp.get('DOB')),
                clean_str_val(emp.get('Father_Name')),
                clean_str_val(emp.get('Bank_Acc_No')),
                clean_str_val(emp.get('Bank_IFSC')),
                clean_str_val(emp.get('UAN_No')),
                clean_str_val(emp.get('ESI_No')),
                clean_str_val(emp.get('Phone_Number')),
                clean_str_val(emp.get('Email_ID')),
                float(emp.get('Fixed_Gross') or 0.0),
                float(emp.get('Basic_DA') or 0.0),
                float(emp.get('HRA') or 0.0),
                float(emp.get('Conveyance_Allowance') or 0.0),
                float(emp.get('Washing_Allowance') or 0.0),
                float(emp.get('Other_Allowance') or 0.0),
                float(emp.get('Per_Day_Wage') or 0.0),
                float(emp.get('OT_Rate') or 0.0),
                pf_val,
                esi_val,
                float(emp.get('LIC') or 0.0),
            ]
            rows_to_write.append(row_data)
    elif include_sample:
        # Guidance sample row for new installations
        sample_row = [
            '1001', 'Rajesh Kumar', 'STAFF', 'PF_ESI', 'Production',
            'CNC Operator', 'Grade A', 'Active', '2024-01-15', '1995-08-20',
            'Ramesh Kumar', '918273645012', 'SBIN0001234', '100918273645',
            '3109182736', '+919876543210', 'rajesh@example.com',
            25000.0, 12500.0, 5000.0, 2500.0, 2500.0, 2500.0, 0.0, 81.25,
            'YES', 'YES', 500.0
        ]
        rows_to_write.append(sample_row)

    # Append rows and apply styling
    for r_idx, row_values in enumerate(rows_to_write, start=2):
        ws.append(row_values)
        ws.row_dimensions[r_idx].height = 21
        is_even = (r_idx % 2 == 0)

        for c_idx, (col_name, dict_key, col_type, _) in enumerate(columns_config, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = regular_font
            cell.border = cell_border
            if not is_even:
                cell.fill = zebra_fill

            if col_type == 'currency':
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif col_type == 'center':
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_type == 'text':
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Set column widths
    for c_idx, (_, _, _, default_w) in enumerate(columns_config, start=1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = default_w

    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_attendance_template_excel(year, month, employees, standard_days=26.0, worker_working_days=None, staff_working_days=None, trans_map=None):
    """
    Generate professional Attendance & Monthly Input Excel workbook pre-filled with live data,
    including Opening Advance, New Advance, Advance Deduction, and Closing Advance.
    Sets separate Company Working Days for Workers vs Staff.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if trans_map is None:
        trans_map = {}

    # Query active advances map from DB to pre-fill opening balances & installments
    adv_map = {}
    try:
        from db import get_db_connection
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("""
            SELECT Emp_No, SUM(Remaining_Amount) AS Total_Remaining, SUM(Monthly_Amount) AS Total_Monthly
            FROM Advances
            WHERE Status = 'Active' AND Remaining_Amount > 0
            GROUP BY Emp_No
        """)
        for row in cur.fetchall():
            emp_no_val = str(row[0]).strip()
            adv_map[emp_no_val] = {
                'remaining': float(row[1] or 0.0),
                'monthly': float(row[2] or 0.0)
            }
        conn.close()
    except Exception as e:
        print(f"[ADVANCE PREFILL NOTICE]: {e}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Attendance_{month}_{year}"
    ws.views.sheetView[0].showGridLines = True

    # Column configuration: (Header name, Dict Key, Type, Default width, is_advance_col)
    columns_config = [
        ('Emp ID', 'Emp_ID', 'text', 12, False),
        ('Employee Name', 'Employee_Name', 'text', 24, False),
        ('Type', 'Type', 'center', 12, False),
        ('Category', 'Category', 'text', 20, False),
        ('Company Working Days', 'Company_Working_Days', 'number', 16, False),
        ('Present', 'Present', 'number', 12, False),
        ('N/H', 'NH', 'number', 10, False),
        ('EL', 'EL', 'number', 10, False),
        ('CL', 'CL', 'number', 10, False),
        ('SL', 'SL', 'number', 10, False),
        ('OT Hours', 'OT_Hours', 'number', 12, False),
        ('Opening Advance', 'Opening_Advance', 'currency', 16, True),
        ('New Advance', 'New_Advance', 'currency', 16, True),
        ('Advance Deduction', 'Advance_Deduction', 'currency', 18, True),
        ('Closing Advance', 'Closing_Advance', 'currency', 16, True),
        ('Arrears', 'Arrears', 'currency', 14, False),
        ('NAPS', 'NAPS', 'currency', 14, False),
        ('LIC', 'LIC', 'currency', 14, False),
        ('Accommodation', 'Accommodation', 'currency', 16, False),
        ('Other', 'Other', 'currency', 14, False)
    ]

    # Styling definitions
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Slate dark
    adv_header_fill = PatternFill(start_color='0E7490', end_color='0E7490', fill_type='solid') # Teal blue for Advance columns

    regular_font = Font(name='Segoe UI', size=9, color='000000')
    bold_font = Font(name='Segoe UI', size=9, bold=True, color='000000')
    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    thin_border = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # Write headers
    for c_idx, (header_name, _, _, _, is_adv) in enumerate(columns_config, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header_name)
        cell.font = header_font
        cell.fill = adv_header_fill if is_adv else header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Write data rows
    for r_idx, emp in enumerate(employees, start=2):
        emp_id = emp['Employee_ID']
        emp_no = str(emp.get('Emp_No', '')).strip()
        existing_t = trans_map.get(emp_id) or {}
        adv_info = adv_map.get(emp_no) or {}

        emp_type_str = str(emp.get('Employee_Type', '')).upper()
        if 'STAFF' in emp_type_str:
            emp_working_days = staff_working_days if staff_working_days is not None else standard_days
        else:
            emp_working_days = worker_working_days if worker_working_days is not None else standard_days

        # Determine Opening Advance
        if 'Opening_Advance' in existing_t and float(existing_t['Opening_Advance'] or 0.0) > 0:
            open_adv = float(existing_t['Opening_Advance'])
        elif adv_info.get('remaining', 0.0) > 0:
            open_adv = adv_info['remaining']
        else:
            open_adv = 0.0

        new_adv = float(existing_t.get('New_Advance', 0.0) or 0.0)

        # Determine Advance Deduction
        if 'Advance_Deduction' in existing_t and float(existing_t['Advance_Deduction'] or 0.0) > 0:
            adv_ded = float(existing_t['Advance_Deduction'])
        elif adv_info.get('monthly', 0.0) > 0:
            adv_ded = min(open_adv + new_adv, adv_info['monthly'])
        else:
            adv_ded = 0.0

        close_adv = max(0.0, open_adv + new_adv - adv_ded)

        row_data = {
            'Emp_ID': emp.get('Emp_No'),
            'Employee_Name': emp.get('Employee_Name'),
            'Type': emp.get('Employee_Type', ''),
            'Category': emp.get('Category', ''),
            'Company_Working_Days': emp_working_days,
            'Present': float(existing_t.get('Present_Days', emp_working_days)),
            'NH': float(existing_t.get('NH', 0.0)),
            'EL': float(existing_t.get('EL', 0.0)),
            'CL': float(existing_t.get('CL', 0.0)),
            'SL': float(existing_t.get('SL', 0.0)),
            'OT_Hours': float(existing_t.get('Act_OT_Hrs', 0.0)),
            'Opening_Advance': open_adv,
            'New_Advance': new_adv,
            'Advance_Deduction': adv_ded,
            'Closing_Advance': close_adv,
            'Arrears': float(existing_t.get('Arrears', 0.0)),
            'NAPS': float(existing_t.get('NAPS_Deduction', 0.0)),
            'LIC': float(existing_t.get('LIC_Deduction', 0.0)),
            'Accommodation': float(existing_t.get('Accommodation_Deduction', 0.0)),
            'Other': float(existing_t.get('Other_Deduction', 0.0))
        }

        is_even = (r_idx % 2 == 0)
        ws.row_dimensions[r_idx].height = 20

        for c_idx, (_, dict_key, col_type, _, is_adv) in enumerate(columns_config, start=1):
            val = row_data.get(dict_key)

            # Excel formula for Closing Advance (Column O, c_idx=15)
            # L=12 (Opening), M=13 (New), N=14 (Deduction), O=15 (Closing)
            if dict_key == 'Closing_Advance':
                cell = ws.cell(row=r_idx, column=c_idx, value=f"=MAX(0, L{r_idx}+M{r_idx}-N{r_idx})")
            else:
                cell = ws.cell(row=r_idx, column=c_idx, value=val)

            cell.font = bold_font if is_adv else regular_font
            cell.border = cell_border
            if not is_even:
                cell.fill = zebra_fill

            if col_type == 'currency':
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif col_type == 'number':
                cell.number_format = '0.0'
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_type == 'center':
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_type == 'text':
                cell.number_format = '@'
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Set column widths
    for c_idx, (_, _, _, default_w, _) in enumerate(columns_config, start=1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = default_w

    ws.freeze_panes = 'C2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

